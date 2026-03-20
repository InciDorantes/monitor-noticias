from bd import obtener_noticias
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from datetime import date
import datetime

fecha_hoy = str(date.today())


def generar_doc(sorted_resultados_filtrados):
    first = 0
    doc = docx.Document()

    #Titulo
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('Noticias Relevantes')
    run.font.size = Pt(14)
    run.bold = True
    run.font.name = 'Lato'

    #Subtitulo
    subtitulo = doc.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitulo.add_run(fecha_hoy)
    run.font.size = Pt(12)
    run.font.name = 'Lato'

    if sorted_resultados_filtrados:
        for noticias in sorted_resultados_filtrados:
            id = noticias['id']
            resumen =  noticias['resumen']
            info_noticia = obtener_noticias(id)
            for info in info_noticia:
                titulo = info[2]
                medio = info[0]
                seccion = info[1]
                fecha_extraida = info[3]
                objeto_fecha = datetime.datetime.strptime(fecha_extraida, '%a, %d %b %Y %H:%M:%S %z')
                fecha = objeto_fecha.strftime('%Y-%m-%d')
            if resumen:
                #nombre de la noticia
                frase = f"{titulo} - {medio}, {seccion}"
                noticia = doc.add_paragraph()
                noticia.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = noticia.add_run(frase)
                run.font.size = Pt(12)
                run.bold = True
                run.font.name = 'Lato'
                noticia.paragraph_format.space_after = Pt(0)
                
                #fecha de publi
                fecha_link = f"{fecha}, {id}"
                fecha_text = doc.add_paragraph()
                fecha_text.paragraph_format.space_before = Pt(0)
                fecha_text.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = fecha_text.add_run(fecha_link)
                run.font.size = Pt(10)
                run.font.name = 'Lato'
                #resumen
                resu = doc.add_paragraph()
                resu.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                run = resu.add_run(resumen)
                run.font.size = Pt(12)
                run.font.name = 'Lato'
            else:
                if first == 0:
                    title_dos = doc.add_paragraph()
                    title_dos.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = title_dos.add_run('Otras Noticias Relevantes')
                    run.font.size = Pt(12)
                    run.bold = True
                    run.font.name = 'Lato'

                    first = 1

                frase = f"{titulo} - {medio} - {seccion} - {fecha} - {id}"
                p = doc.add_paragraph()
                p.style = 'List Bullet'
                r = p.add_run(frase)
                r.font.name = 'Lato'
                p.paragraph_format.line_spacing = Pt(18)

    doc.save('noticias.docx')